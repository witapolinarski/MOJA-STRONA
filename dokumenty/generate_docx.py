#!/usr/bin/env python3
"""Generuje wersje .docx statutu na podstawie statut-stowarzyszenia-strzeleckiego-wzor.md.

Uzycie: python3 generate_docx.py
Wymaga: pip install python-docx
"""
import re
from pathlib import Path

from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH

SRC = Path(__file__).parent / "statut-stowarzyszenia-strzeleckiego-wzor.md"
OUT = Path(__file__).parent / "statut-stowarzyszenia-strzeleckiego-wzor.docx"


def set_base_style(doc):
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)
    for section in doc.sections:
        section.top_margin = Cm(2)
        section.bottom_margin = Cm(2)
        section.left_margin = Cm(2.5)
        section.right_margin = Cm(2.5)


def add_title(doc, text, size=16, bold=True, align_center=True):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.bold = bold
    run.font.size = Pt(size)
    if align_center:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    return p


def add_heading(doc, text):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.bold = True
    run.font.size = Pt(13)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.space_before = Pt(12)
    return p


def add_paragraph_section(doc, text):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.bold = True
    return p


def add_body(doc, text):
    p = doc.add_paragraph(text)
    return p


def strip_md_inline(text):
    text = re.sub(r"\*\*(.*?)\*\*", r"\1", text)
    text = re.sub(r"(?<!\*)\*([^*]+)\*(?!\*)", r"\1", text)
    text = re.sub(r"`(.*?)`", r"\1", text)
    return text


def main():
    lines = SRC.read_text(encoding="utf-8").splitlines()
    doc = Document()
    set_base_style(doc)

    i = 0
    n = len(lines)
    in_table = False
    table = None

    while i < n:
        raw = lines[i]
        line = raw.rstrip()
        stripped = line.strip()

        if stripped.startswith("# "):
            add_title(doc, strip_md_inline(stripped[2:]).strip(), size=18)
            i += 1
            continue
        if stripped.startswith("## "):
            add_heading(doc, strip_md_inline(stripped[3:]).strip())
            i += 1
            continue
        if stripped.startswith("### "):
            add_paragraph_section(doc, strip_md_inline(stripped[4:]).strip())
            i += 1
            continue
        if stripped == "---":
            i += 1
            continue
        if stripped.startswith("> "):
            p = doc.add_paragraph()
            run = p.add_run(strip_md_inline(stripped[2:]).strip())
            run.italic = True
            run.font.size = Pt(9.5)
            i += 1
            continue
        if stripped.startswith("**§"):
            text = strip_md_inline(stripped)
            p = doc.add_paragraph()
            run = p.add_run(text)
            run.bold = True
            p.space_before = Pt(8)
            i += 1
            continue
        if stripped.startswith("|"):
            # collect table block
            table_lines = []
            while i < n and lines[i].strip().startswith("|"):
                table_lines.append(lines[i].strip())
                i += 1
            rows = [l.strip("|").split("|") for l in table_lines]
            rows = [[c.strip() for c in r] for r in rows]
            # drop separator row (----)
            rows = [r for r in rows if not all(set(c) <= set("-: ") for c in r)]
            if rows:
                t = doc.add_table(rows=len(rows), cols=len(rows[0]))
                t.style = "Table Grid"
                for r_idx, row in enumerate(rows):
                    for c_idx, cell in enumerate(row):
                        t.cell(r_idx, c_idx).text = strip_md_inline(cell)
            continue
        if stripped == "":
            i += 1
            continue
        if stripped.startswith("- "):
            text = strip_md_inline(stripped[2:])
            p = doc.add_paragraph(f"•  {text}")
            p.paragraph_format.left_indent = Cm(0.6)
            i += 1
            continue

        # numbered / bulleted sub-items
        m_num = re.match(r"^(\s*)(\d+)\.\s+(.*)$", line)
        if m_num:
            indent = len(m_num.group(1))
            text = strip_md_inline(m_num.group(3))
            p = doc.add_paragraph(f"{m_num.group(2)}. {text}")
            p.paragraph_format.left_indent = Cm(0.6 + (0.6 if indent > 0 else 0))
            i += 1
            continue

        add_body(doc, strip_md_inline(stripped))
        i += 1

    doc.save(OUT)
    print(f"Zapisano: {OUT}")


if __name__ == "__main__":
    main()
